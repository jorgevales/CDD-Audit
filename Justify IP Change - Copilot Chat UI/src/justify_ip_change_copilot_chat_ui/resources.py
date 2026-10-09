from __future__ import annotations

from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from .errors import ResourceError
from .models import CaseRecord


REQUIRED_CASE_COLUMNS = (
    "change_id",
    "InterestedPartyId",
    "InterestedPartyCurrentName",
    "Date_of_birth",
    "Status",
    "ActionDateTime",
    "ActionUserId",
    "ActionUserName",
    "ActionUserTeam",
    "ChangedSections",
    "ChangedFields",
    "PreviousValues",
    "NewValues",
    "FieldChangeCount",
)


def load_text_resource(path: Path) -> str:
    try:
        value = path.read_text(encoding="utf-8-sig")
    except (OSError, UnicodeError) as exc:
        raise ResourceError(f"Could not read runtime resource {path.name}: {exc}") from exc
    if not value.strip():
        raise ResourceError(f"Runtime resource is empty: {path.name}")
    return value


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def load_case_workbook(path: Path) -> list[CaseRecord]:
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise ResourceError(f"Could not open the runtime workbook {path.name}: {exc}") from exc
    matching = []
    try:
        for sheet in workbook.worksheets:
            iterator = sheet.iter_rows(values_only=True)
            header = next(iterator, None)
            if not header:
                continue
            names = tuple(_cell_text(value) for value in header)
            if all(name in names for name in REQUIRED_CASE_COLUMNS):
                matching.append((sheet.title, names, list(iterator)))
        if len(matching) != 1:
            raise ResourceError(
                f"Workbook must contain exactly one worksheet with the required Step 08 columns; found {len(matching)}."
            )
        _, names, raw_rows = matching[0]
        records: list[CaseRecord] = []
        seen: set[tuple[str, str]] = set()
        for number, values in enumerate(raw_rows, 2):
            row = {name: _cell_text(values[index] if index < len(values) else None) for index, name in enumerate(names)}
            change_id = row.get("change_id", "")
            party_id = row.get("InterestedPartyId", "")
            if not change_id and not party_id and not any(row.values()):
                continue
            if not change_id or not party_id:
                raise ResourceError(f"Workbook row {number} has a blank change_id or InterestedPartyId.")
            record = CaseRecord(change_id, party_id, row)
            if record.canonical_key in seen:
                raise ResourceError(f"Workbook contains a duplicate case at row {number}: {change_id}/{party_id}.")
            seen.add(record.canonical_key)
            records.append(record)
        if not 1 <= len(records) <= 1000:
            raise ResourceError(f"Workbook must contain 1 to 1,000 case rows; found {len(records)}.")
        return records
    finally:
        workbook.close()


def case_records_in_range(records: Iterable[CaseRecord], low: int, high: int) -> list[CaseRecord]:
    selected = []
    for record in records:
        if record.change_id.strip().isdigit() and low <= int(record.change_id) <= high:
            selected.append(record)
    return sorted(selected, key=lambda item: (int(item.change_id), item.interested_party_id.casefold()))
